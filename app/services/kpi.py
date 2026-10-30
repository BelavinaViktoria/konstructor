"""Разбор карты KPI (.xlsx). Суммы премий, ставки и табельный номер в план не переносятся."""
import re
from datetime import date

from openpyxl import load_workbook

from .dates import add_months, fmt_date, prev_workday

MAX_ROWS = 2000
MAX_COLS = 40
MONEY = re.compile(r"[=,.;]?\s*\d[\d\s]*\s*(руб[а-яё]*|рбл[а-яё]*|р\.|₽)\.?", re.I)


def read_sheets(stream) -> list[dict]:
    """Читает книгу в вид [{'name':..., 'g': {строка: {колонка: текст}}}]."""
    wb = load_workbook(stream, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            g: dict[int, dict[int, str]] = {}
            for r, row in enumerate(ws.iter_rows(min_row=1, max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True), start=1):
                for c, value in enumerate(row, start=1):
                    if value is not None and str(value).strip():
                        g.setdefault(r, {})[c] = str(value).strip()
            sheets.append({"name": ws.title, "g": g})
    finally:
        wb.close()
    return sheets


def parse_sheet(sheet: dict) -> dict | None:
    g = sheet["g"]
    rows = sorted(g)

    def cell(r: int, c: int) -> str:
        return g.get(r, {}).get(c, "")

    header = next((r for r in rows if any(re.search(r"показател|кпэ|kpi", v, re.I) for v in g[r].values())), None)
    if header is None or not any(re.search(r"периодичн|услови", v, re.I) for v in g[header].values()):
        return None

    def col_of(pattern: str) -> int:
        for r in (header, header + 1):
            for c, v in sorted(g.get(r, {}).items()):
                if re.search(pattern, v, re.I):
                    return c
        return 0

    c_name = col_of(r"^наименовани") or col_of(r"показател|кпэ|kpi")
    c_per, c_cond = col_of(r"периодичн"), col_of(r"услови")
    resp_cols = [c for c, v in sorted(g[header].items()) if re.search("ответствен", v, re.I)]
    c_resp = next((c for c in resp_cols if any(
        r > header and re.search(r"руководител|[А-ЯЁ]\.\s?[А-ЯЁ]\.", cell(r, c)) and len(cell(r, c)) < 90 for r in rows)), 0)

    before = [" ".join(g[r].values()) for r in rows if r < header]
    who = next((v for v in before if re.search(r"таб\.?\s*номер|,", v, re.I)), "")
    title = next((v for v in before if re.search(r"выплата|период|за\s+[а-я]+", v, re.I)), "")
    parts = [x.strip() for x in who.split(",") if x.strip() and not re.search("таб", x, re.I)]

    indicators, cur = [], None
    after = [r for r in rows if r > header]
    for r in after:
        name = cell(r, c_name)
        if name and len(name) > 2 and not re.search(r"^наименовани", name, re.I) and (cell(r, c_per) if c_per else True):
            cur = {"name": re.sub(r"\s+", " ", name).strip(), "per": cell(r, c_per), "cond": [], "resp": cell(r, c_resp), "row": r}
            indicators.append(cur)
        elif cur and not name and cell(r, c_cond) and not cell(r, 1):
            cur["cond"].append(cell(r, c_cond))
        if cur and cur["row"] == r and cell(r, c_cond):
            cur["cond"].insert(0, cell(r, c_cond))
        if any(re.search(r"подпись|сотрудник:|согласовано", v, re.I) for v in g[r].values()):
            cur = None

    smart = []
    for r in after:
        m = re.match(r"^\s*(\d+)[.)]\s+(.{6,})", cell(r, 1))
        if m:
            smart.append({"name": m.group(2).strip(), "row": r})
    return {"sheet": sheet["name"], "fio": parts[0] if parts else "", "pos": parts[-1] if len(parts) > 1 else "",
            "period": title, "K": indicators, "smart": smart}


def criterion(cond: list[str]) -> str:
    text = "; ".join(cond)
    if re.search(r"К\s*=", text):
        return "Оценка по сроку: в срок К 1,0; с нарушением срока К 0,8; не выполнена К 0"
    clean = []
    for x in cond:
        y = re.sub(r"без потолка", "", MONEY.sub("", x), flags=re.I)
        y = re.sub(r"\s*=\s*", " ", y).strip(" .,;")
        if y:
            clean.append(y)
    if not clean:
        return "Критерий в карте не указан: уточните у руководителя"
    has_money = any(MONEY.search(x) for x in cond)
    return "Результат считается по количеству: " + "; ".join(clean) if has_money else "; ".join(clean)


def suggest_due(name: str, hire: date | None, today: date | None = None) -> tuple[date, int]:
    """Срок освоения по смыслу задачи, в пределах трёх месяцев от приёма."""
    if re.search(r"освоени|освоить|crm|регламент|аттестац|стандарт|скрипт|ведение|обучени|изуч", name, re.I):
        months = 1
    elif re.search(r"встреч|привлеч|подготов|предложени|клиент|резерв|студент|курс", name, re.I):
        months = 2
    else:
        months = 3
    base = hire or today or date.today()
    return prev_workday(add_months(base, months)), months


def build_items(parsed: dict, hire: date | None) -> list[dict]:
    """Пункты карты, готовые стать целями; алгоритм и формат нужны при переносе в задачи."""
    items = []
    smart_ind = next((k for k in parsed["K"] if re.search("смарт|smart", k["name"], re.I)), None)
    for k in parsed["K"]:
        if smart_ind is k and parsed["smart"]:
            continue
        crit = criterion(k["cond"])
        if k is smart_ind:
            crit += ". Сами SMART-задачи в карте не перечислены: руководитель уточнит их при согласовании плана"
        due, months = suggest_due(k["name"], hire)
        resp = "Сотрудник"
        if k["resp"]:
            resp += "; данные: " + re.sub(r"Руководитель\s+", "руководитель ", k["resp"], flags=re.I)
        items.append({
            "title": k["name"][:1].upper() + k["name"][1:], "criterion": crit, "fmt": "KPI-показатель",
            "due": due.isoformat(), "months": months, "responsible": resp,
            "source": f"лист «{parsed['sheet']}», строка {k['row']}",
            "algorithm": [
                "Уточните у руководителя, как измеряется показатель: " + crit[:1].lower() + crit[1:],
                f"Вместе с наставником разберите, как выполняется эта функция, и составьте план освоения до срока {fmt_date(due)}",
                "Соберите подтверждение результата: отчёт, выгрузку или ссылку на документ",
                "Отметьте выполнение в плане, приложите подтверждение и комментарий для руководителя",
            ],
        })
    for x in parsed["smart"]:
        due, months = suggest_due(x["name"], hire)
        crit = criterion(smart_ind["cond"] if smart_ind else [])
        items.append({
            "title": re.sub(r"\s+", " ", x["name"]), "criterion": "SMART-задача из карты KPI. " + crit, "fmt": "SMART-задача",
            "due": due.isoformat(), "months": months, "responsible": "Сотрудник",
            "source": f"лист «{parsed['sheet']}», строка {x['row']}",
            "algorithm": [
                "Уточните с руководителем ожидаемый результат и срок задачи", "Разбейте задачу на шаги и согласуйте их с наставником",
                "Выполните задачу и соберите подтверждение результата", "Отметьте выполнение, приложите ссылку и комментарий",
            ],
        })
    return items


def parse_workbook(sheets: list[dict], hire: date | None) -> dict | None:
    """Берёт первый лист с показателями. Ничего не найдено: None."""
    for sheet in sheets:
        parsed = parse_sheet(sheet)
        if parsed and (parsed["K"] or parsed["smart"]):
            return {"fio": parsed["fio"], "pos": parsed["pos"], "period": parsed["period"], "sheet": parsed["sheet"],
                    "items": build_items(parsed, hire)}
    return None
